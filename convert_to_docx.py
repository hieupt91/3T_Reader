import os
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def md_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip('\n')
        if not line:
            doc.add_paragraph()
            continue
            
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        else:
            p = doc.add_paragraph()
            if line.startswith('- '):
                p.style = 'List Bullet'
                line = line[2:]
            
            # Simple bold parsing **text**
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    # simple italic parsing *text*
                    iparts = re.split(r'(\*.*?\*)', part)
                    for ipart in iparts:
                        if ipart.startswith('*') and ipart.endswith('*'):
                            run = p.add_run(ipart[1:-1])
                            run.italic = True
                        else:
                            p.add_run(ipart)
                            
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    doc.save(docx_path)
    print(f"Saved {docx_path}")

base_dir = r"docs\Ho_So_Phap_Ly"
files = ["TO_KHAI_BAN_QUYEN_TAC_GIA.md", "TO_KHAI_NHAN_HIEU.md", "BAN_MO_TA_PHAN_MEM.md", "EVALUATION_BAN_QUYEN.md", "HUONG_DAN_NOP_HO_SO.md"]

for file in files:
    md_path = os.path.join(base_dir, file)
    docx_path = os.path.join(base_dir, file.replace('.md', '.docx'))
    if os.path.exists(md_path):
        md_to_docx(md_path, docx_path)
