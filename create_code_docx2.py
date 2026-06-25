# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_source_code_docx():
    src_file = r"docs\Ho_So_Phap_Ly\SOURCE_CODE_COPYRIGHT.txt"
    out_file = r"docs\Ho_So_Phap_Ly\MA_NGUON_50_TRANG_DAU_VA_CUOI.docx"
    
    if not os.path.exists(src_file):
        print("Source text not found.")
        return
        
    with open(src_file, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()
        
    first_lines = lines[:3000]
    last_lines = lines[-3000:] if len(lines) > 6000 else []
    
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Consolas'
    font.size = Pt(9)
    
    p = doc.add_paragraph("MÃ NGUỒN PHẦN MỀM 3T READER\n(50 TRANG ĐẦU VÀ 50 TRANG CUỐI)")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.name = 'Times New Roman'
    doc.add_page_break()
    
    p = doc.add_paragraph("--- 50 TRANG ĐẦU TIÊN ---\n")
    p.runs[0].bold = True
    doc.add_paragraph("".join(first_lines))
    
    if last_lines:
        doc.add_page_break()
        p = doc.add_paragraph("--- 50 TRANG CUỐI CÙNG ---\n")
        p.runs[0].bold = True
        doc.add_paragraph("".join(last_lines))
        
    doc.save(out_file)
    print("Created MA_NGUON_50_TRANG_DAU_VA_CUOI.docx")

if __name__ == "__main__":
    create_source_code_docx()
