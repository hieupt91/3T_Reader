import json
import os
from datetime import datetime

from packages.qt_compat.QtWidgets import QFileDialog

from app.actions._guard import require_document
from app.dialogs import show_info, show_warning


def _format_size(size_bytes: int) -> str:
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    if size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.2f} MB"
    return f"{size_bytes / (1024 * 1024 * 1024):.2f} GB"


def _run_find(window, query: str, *, find_previous: bool, new_search: bool):
    wv = window._get_webview()
    if not wv:
        show_warning(window, "Chua san sang", "Trinh xem PDF chua san sang de tim kiem.")
        return

    payload = {
        "query": query,
        "phraseSearch": True,
        "caseSensitive": False,
        "entireWord": False,
        "highlightAll": True,
        "findPrevious": find_previous,
        "matchDiacritics": False,
        "type": "" if new_search else "again",
    }

    js = f"""
(function() {{
    try {{
        const app = window.PDFViewerApplication;
        if (!app || !app.eventBus) {{
            return "PDF.js chua khoi tao xong";
        }}
        app.eventBus.dispatch("find", Object.assign({{ source: window }}, {json.dumps(payload)}));
        return "OK";
    }} catch (e) {{
        return String(e);
    }}
}})();
"""

    def _after(result):
        if result == "OK":
            window.status.showMessage(
                f"Dang tim: '{query}'" + (" (lui)" if find_previous else ""),
                3000,
            )
            viewer = window.viewer
            if viewer and new_search:
                viewer.check_find_result(query, delay_ms=700)
        elif result:
            show_warning(window, "Khong the tim kiem", str(result))

    wv.page().runJavaScript(js, _after)


def execute_search(window, query: str, *, find_previous: bool = False, new_search: bool = True) -> bool:
    query = (query or "").strip()
    if not query:
        return False
    window.search_query = query
    _run_find(window, query, find_previous=find_previous, new_search=new_search)
    return True


@require_document(show_message=True)
def search_text(window):
    if hasattr(window, "show_search_panel"):
        window.show_search_panel()


def _search_direction(window, *, find_previous: bool):
    if not window.current_path:
        show_warning(window, "Chua mo tep", "Vui long mo tep PDF truoc khi tim kiem.")
        return

    query = getattr(window, "search_query", "").strip()
    if not query:
        search_text(window)
        return
    execute_search(window, query, find_previous=find_previous, new_search=False)


def search_next(window):
    _search_direction(window, find_previous=False)


def search_previous(window):
    _search_direction(window, find_previous=True)


@require_document(show_message=True)
def show_file_info(window):
    source_path = window.current_path
    display_path = window.get_display_path() if hasattr(window, "get_display_path") else source_path
    try:
        stat = os.stat(source_path)
        size_text = _format_size(stat.st_size)
        modified_text = datetime.fromtimestamp(stat.st_mtime).strftime("%d/%m/%Y %H:%M:%S")
    except OSError as exc:
        show_warning(window, "Khong doc duoc thong tin tep", str(exc))
        return

    page_count = window.viewer.get_page_count() if source_path else 0
    current_page = window.viewer.get_current_page() if source_path else 0

    message = (
        f"Ten tep: {os.path.basename(display_path)}\n"
        f"Duong dan: {display_path}\n"
        f"Dung luong: {size_text}\n"
        f"So trang: {page_count}\n"
        f"Trang hien tai: {current_page}\n"
        f"Cap nhat lan cuoi: {modified_text}"
    )
    show_info(window, "Thong tin tep PDF", message)


def _pick_export_path(window, title: str, default_name: str, file_filter: str, extension: str) -> str | None:
    path, _ = QFileDialog.getSaveFileName(window, title, default_name, file_filter)
    if not path:
        return None
    if not path.lower().endswith(extension):
        path += extension
    return path


def _active_pdf_source(window) -> str | None:
    return getattr(window, "current_path", None)


@require_document(show_message=True)
def export_to_docx(window):
    source_path = _active_pdf_source(window)
    if not source_path or not os.path.exists(source_path):
        show_warning(window, "Khong the xuat DOCX", "Khong tim thay tep PDF hien tai.")
        return

    default_name = os.path.splitext(os.path.basename(source_path))[0] + ".docx"
    target_path = _pick_export_path(
        window,
        "Xuat sang DOCX",
        default_name,
        "Word Document (*.docx)",
        ".docx",
    )
    if not target_path:
        return

    try:
        from pdf2docx import Converter
    except ImportError:
        show_warning(window, "Thieu dependency DOCX", "Chua cai `pdf2docx`. Hay cai dependency roi thu lai.")
        return

    converter = None
    try:
        converter = Converter(source_path)
        converter.convert(target_path)
        window.status.showMessage(f"Da xuat DOCX: {os.path.basename(target_path)}", 4000)
    except Exception as exc:
        show_warning(window, "Xuat DOCX that bai", str(exc))
    finally:
        if converter is not None:
            try:
                converter.close()
            except Exception:
                pass


@require_document(show_message=True)
def export_to_excel(window):
    source_path = _active_pdf_source(window)
    if not source_path or not os.path.exists(source_path):
        show_warning(window, "Khong the xuat Excel", "Khong tim thay tep PDF hien tai.")
        return

    default_name = os.path.splitext(os.path.basename(source_path))[0] + ".xlsx"
    target_path = _pick_export_path(
        window,
        "Xuat sang Excel",
        default_name,
        "Excel Workbook (*.xlsx)",
        ".xlsx",
    )
    if not target_path:
        return

    try:
        import pdfplumber
        from openpyxl import Workbook
    except ImportError:
        show_warning(window, "Thieu dependency Excel", "Chua cai `pdfplumber` hoac `openpyxl`. Hay cai dependency roi thu lai.")
        return

    try:
        workbook = Workbook()
        first_sheet = workbook.active
        workbook.remove(first_sheet)

        sheet_count = 0
        with pdfplumber.open(source_path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                for table_index, table in enumerate(tables, start=1):
                    sheet = workbook.create_sheet(title=f"P{page_index}_T{table_index}")
                    sheet_count += 1
                    for row in table or []:
                        sheet.append([cell if cell is not None else "" for cell in row])

            if sheet_count == 0:
                sheet = workbook.create_sheet(title="PDF_Text")
                for page_index, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    sheet.append([f"Page {page_index}"])
                    for line in text.splitlines():
                        sheet.append([line])
                    sheet.append([""])

        workbook.save(target_path)
        window.status.showMessage(f"Da xuat Excel: {os.path.basename(target_path)}", 4000)
    except Exception as exc:
        show_warning(window, "Xuat Excel that bai", str(exc))
