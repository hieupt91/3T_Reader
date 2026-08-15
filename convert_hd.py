# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import re

def md_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip('\n')
        if not line:
            doc.add_paragraph()
            continue
            
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        else:
            p = doc.add_paragraph()
            if line.startswith('- '):
                p.style = 'List Bullet'
                line = line[2:]
            elif re.match(r'^\d+\.\s', line):
                p.style = 'List Number'
                line = line.split(' ', 1)[1]
            
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    iparts = re.split(r'(\*.*?\*)', part)
                    for ipart in iparts:
                        if ipart.startswith('*') and ipart.endswith('*'):
                            run = p.add_run(ipart[1:-1])
                            run.italic = True
                        else:
                            p.add_run(ipart)
                            
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    doc.save(docx_path)
    print(f"Saved {docx_path}")

md_to_docx(r"docs\Ho_So_Phap_Ly\HUONG_DAN_SU_DUNG_PHAN_MEM.md", r"docs\Ho_So_Phap_Ly\HUONG_DAN_SU_DUNG_PHAN_MEM.docx")
