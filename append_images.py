# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def append_images_to_docx(docx_path, images, title):
    try:
        doc = Document(docx_path)
    except Exception as e:
        print(f"Error opening {docx_path}: {e}")
        return

    doc.add_page_break()
    p_title = doc.add_heading(title, level=1)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for img_path, caption in images:
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            try:
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.0))
                p_cap = doc.add_paragraph(caption)
                p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                p_cap.style.font.italic = True
                p_cap.style.font.size = Pt(11)
                doc.add_paragraph() # spacing
                print(f"Added {img_path}")
            except Exception as e:
                print(f"Failed to add {img_path}: {e}")
        else:
            print(f"Missing image {img_path}")
            
    doc.save(docx_path)
    print(f"Saved {docx_path}")

images_to_add = [
    (r"assets\brand_appicon_512.png", "Hình 1: Logo và Nhãn hiệu 3T Reader"),
    (r"assets\brand_banner.png", "Hình 2: Giao diện và Banner giới thiệu phần mềm"),
    (r"assets\dmg_background.jpg", "Hình 3: Giao diện Cài đặt phần mềm trên macOS"),
    (r"assets\feat_secure.png", "Hình 4: Tính năng Ký số điện tử bảo mật (e-Sign)")
]

append_images_to_docx(r"docs\Ho_So_Phap_Ly\EVALUATION_BAN_QUYEN.docx", images_to_add, "PHỤ LỤC HÌNH ẢNH NHÃN HIỆU VÀ GIAO DIỆN")
append_images_to_docx(r"docs\Ho_So_Phap_Ly\BAN_MO_TA_PHAN_MEM.docx", images_to_add, "PHỤ LỤC HÌNH ẢNH GIAO DIỆN")
