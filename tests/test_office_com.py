from __future__ import annotations

import os
import sys

import pytest

from packages.document_core.office_com import (
    _com_progid_exists,
    _kind_for_extension,
    convert_via_office_com,
    detect_office_suite,
)

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="COM Automation chỉ có trên Windows")


@pytest.mark.parametrize(
    "ext,expected",
    [
        (".doc", "word"), (".docx", "word"), (".rtf", "word"), (".odt", "word"),
        (".xls", "excel"), (".xlsx", "excel"), (".ods", "excel"), (".csv", "excel"),
        (".ppt", "ppt"), (".pptx", "ppt"), (".odp", "ppt"),
        (".pdf", None), (".txt", None), ("", None),
    ],
)
def test_kind_for_extension(ext, expected):
    assert _kind_for_extension(ext) == expected


def test_com_progid_exists_false_for_bogus_progid():
    assert _com_progid_exists("ThisProgIdDefinitelyDoesNotExist.Bogus12345") is False


def test_com_progid_exists_false_for_empty_string():
    assert _com_progid_exists("") is False


def test_convert_via_office_com_returns_false_for_unsupported_extension(tmp_path):
    src = tmp_path / "file.txt"
    src.write_text("not an office file")
    assert convert_via_office_com(str(src), str(tmp_path / "out.pdf")) is False


_OFFICE_SUITE = detect_office_suite() if sys.platform == "win32" else None
_requires_office = pytest.mark.skipif(
    _OFFICE_SUITE is None, reason="Không phát hiện MS Office/WPS đã cài trên máy này"
)


@_requires_office
def test_detect_office_suite_matches_actual_installed_suite():
    """Chạy thật (không mock) trên máy đang có Office cài sẵn - xác nhận
    detect_office_suite() không lỗi và trả 1 trong 2 giá trị hợp lệ."""
    assert _OFFICE_SUITE in {"msoffice", "wps"}


@_requires_office
def test_convert_word_to_pdf_via_real_com(tmp_path):
    """Test hành vi thật đầu-cuối: tạo 1 file .docx thật bằng python-docx,
    convert bằng COM Automation thật (Word/WPS đang cài trên máy), xác nhận
    PDF ra đúng nội dung bằng cách đọc lại qua pypdfium2 - không mock bất kỳ
    lớp nào, vì đây chính là hành vi thật user sẽ gặp."""
    import docx
    import pypdfium2 as pdfium

    src = tmp_path / "b52_word_test.docx"
    out = tmp_path / "b52_word_test.pdf"
    doc = docx.Document()
    doc.add_paragraph("B52 automated test paragraph - COM Automation real conversion.")
    doc.save(str(src))

    ok = convert_via_office_com(str(src), str(out))

    assert ok is True
    assert out.exists() and out.stat().st_size > 100
    pdf = pdfium.PdfDocument(str(out))
    try:
        text = pdf[0].get_textpage().get_text_range()
        assert "COM Automation real conversion" in text
    finally:
        pdf.close()


@_requires_office
def test_convert_excel_to_pdf_via_real_com(tmp_path):
    import openpyxl
    import pypdfium2 as pdfium

    src = tmp_path / "b52_excel_test.xlsx"
    out = tmp_path / "b52_excel_test.pdf"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "B52 excel automated test cell"
    wb.save(str(src))

    ok = convert_via_office_com(str(src), str(out))

    assert ok is True
    assert out.exists() and out.stat().st_size > 100
    pdf = pdfium.PdfDocument(str(out))
    try:
        text = pdf[0].get_textpage().get_text_range()
        assert "B52 excel automated test cell" in text
    finally:
        pdf.close()


@_requires_office
def test_convert_leaves_no_zombie_office_process(tmp_path):
    """Rủi ro đã lường trước trong ROADMAP_B52...: nếu Quit()/CoUninitialize
    không chạy đúng lúc lỗi, WINWORD.EXE/EXCEL.EXE có thể treo lại trong Task
    Manager dù convert đã xong. Test bằng cách đếm process trước/sau."""
    import subprocess
    import docx

    def _office_process_count() -> int:
        names = "WINWORD.EXE,EXCEL.EXE,POWERPNT.EXE"
        result = subprocess.run(
            ["tasklist", "/FI", f"IMAGENAME eq WINWORD.EXE"],
            capture_output=True, text=True, timeout=10,
        )
        count = 0
        for name in names.split(","):
            r = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {name}"], capture_output=True, text=True, timeout=10)
            count += r.stdout.count(".EXE")
        return count

    before = _office_process_count()

    src = tmp_path / "b52_zombie_test.docx"
    out = tmp_path / "b52_zombie_test.pdf"
    doc = docx.Document()
    doc.add_paragraph("zombie process check")
    doc.save(str(src))
    convert_via_office_com(str(src), str(out))

    after = _office_process_count()
    assert after <= before, f"Office process leaked: before={before} after={after}"
